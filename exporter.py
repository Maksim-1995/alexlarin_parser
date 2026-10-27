"""Экспорт заданий в Excel и ZIP-архив с изображениями."""

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def write_excel(tasks, metadata, excel_path):
    """Сохранить задания и метаданные в книгу Excel.

    Лист ``Tasks`` содержит поля ``task_num``, ``condition``,
    ``image_name``, ``solution``, ``answer``, ``exams_id`` и ``topics_id``.
    Поле ``solution``
    остаётся пустым. На листе ``about`` первая строка содержит названия
    полей метаданных, вторая — их значения. Для поля ``date`` задаётся
    формат ``yyyy-mm-dd``.

    Args:
        tasks: Итерируемая коллекция словарей с полями ``task_num``,
            ``condition``, ``image_name``, ``answer``, ``exams_id`` и ``topics_id``.
        metadata: Словарь метаданных для листа ``about``.
        excel_path: Путь к сохраняемой книге Excel.

    Raises:
        KeyError: В задании отсутствует обязательное поле.
        OSError: Не удалось записать книгу по указанному пути.
    """
    workbook = Workbook()

    # При создании книги один лист уже существует.
    tasks_sheet = workbook.active
    tasks_sheet.title = "Tasks"

    tasks_sheet.append([
        "task_num",
        "condition",
        "image_name",
        "solution",
        "answer",
        "exams_id",
        "topics_id",
    ])

    for task in tasks:
        tasks_sheet.append([
            task["task_num"],
            task["condition"],
            task["image_name"],
            "",  # Решения пока оставляем пустыми.
            task["answer"],
            task["exams_id"],
            task["topics_id"],
        ])

    about_sheet = workbook.create_sheet("about")
    about_sheet.append(list(metadata.keys()))
    about_sheet.append(list(metadata.values()))

    for column, field in enumerate(metadata, start=1):
        if field == "date":
            date_cell = about_sheet.cell(
                row=2,
                column=column,
            )
            date_cell.number_format = "yyyy-mm-dd"

    format_excel(workbook)
    workbook.save(excel_path)


def format_excel(workbook):
    """Оформить листы книги Excel, изменяя её на месте.

    На всех листах закрепляется строка заголовков, включается перенос
    текста и задаётся стиль заголовков. На листах ``Tasks`` и ``about``
    настраивается ширина столбцов; для ``Tasks`` включается автофильтр.

    Args:
        workbook: Книга openpyxl с листами ``Tasks`` и ``about``.

    Raises:
        KeyError: В книге отсутствует лист ``Tasks`` или ``about``.
    """
    # Общие настройки обоих листов.
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="305496",
            )

        sheet.row_dimensions[1].height = 25

    tasks_sheet = workbook["Tasks"]

    column_widths = {
        "A": 12,  # task_num
        "B": 90,  # condition
        "C": 30,  # image_name
        "D": 30,  # solution
        "E": 50,  # answer
        "F": 14,  # exams_id
        "G": 14,  # topics_id
    }

    for column, width in column_widths.items():
        tasks_sheet.column_dimensions[column].width = width

    tasks_sheet.auto_filter.ref = tasks_sheet.dimensions

    about_sheet = workbook["about"]
    about_widths = {
        "class": 10,
        "year": 12,
        "date": 16,
        "variant_name": 45,
        "title": 45,
        "source_url": 70,
        "answers_13_19_source": 70,
        "conditions_source": 70,
        "formula_notation": 70,
    }
    for cell in about_sheet[1]:
        about_sheet.column_dimensions[cell.column_letter].width = (
            about_widths.get(cell.value, 32)
        )
    about_sheet.row_dimensions[2].height = 90


def create_zip(excel_path, images_dir, archive_path):
    """Создать ZIP-архив и проверить целостность записанных файлов.

    Книга сохраняется в архиве как ``tasks.xlsx``. Рисунки добавляются
    рядом с ней, без вложенной папки. ``images_dir`` может совпадать
    с папкой Excel: книга повторно не добавляется. Сам архив также
    исключается из списка исходных файлов. Вложенные каталоги
    не обходятся. Существующий архив перезаписывается.

    Каталог изображений должен содержать только итоговые рисунки
    и, при совпадении с папкой результата, книгу Excel.

    Args:
        excel_path: Путь к книге Excel, добавляемой в архив.
        images_dir: Каталог изображений; теперь это папка результата.
        archive_path: Путь к создаваемому ZIP-архиву.

    Returns:
        int: Количество файлов в архиве, включая книгу Excel.

    Raises:
        OSError: Не удалось прочитать исходные файлы или записать архив.
        ValueError: Проверка целостности обнаружила повреждённый файл.
    """
    excel_path = Path(excel_path)
    images_dir = Path(images_dir)
    archive_path = Path(archive_path)
    excluded_paths = {excel_path.resolve(), archive_path.resolve()}

    with ZipFile(
        archive_path,
        mode="w",
        compression=ZIP_DEFLATED,
    ) as archive:
        archive.write(excel_path, arcname="tasks.xlsx")

        for image_path in sorted(images_dir.iterdir()):
            if image_path.is_file() and image_path.resolve() not in excluded_paths:
                archive.write(
                    image_path,
                    arcname=image_path.name,
                )

    # Проверяем целостность записанных файлов.
    with ZipFile(archive_path, mode="r") as archive:
        damaged_file = archive.testzip()

        if damaged_file is not None:
            raise ValueError(
                f"В архиве повреждён файл: {damaged_file}"
            )

        return len(archive.namelist())

"""Сборка реального варианта из input/; классификация в тестах подменяется.

Результаты создаются исключительно во временной папке pytest. Если сохранённые
исходники не предоставлены вместе с клоном, интеграционные проверки пропускаются.
"""

import csv
import json
import shutil
import socket
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

import classification
import downloader
import main
import recognizer
from export_metadata import ABOUT_FIELDS


PROJECT = Path(__file__).resolve().parents[1]
SOURCE_NAMES = (
    "trvar542.html", "trvar542.pdf", "trvar542.js", "trvar542.png",
    "trvar542_answers.json", "trvar542_conditions_gemini.json",
)
IMAGE_NAMES = {1: "image_82.jpeg", 3: "image_83.jpeg", 8: "image_198.png", 11: "image_200.png"}
EXAM_LABEL = "E | Вектор > Векторы и операции с ними"
TOPIC_LABEL = (
    "T | Геометрия > Планиметрия > Векторы > "
    "Умножение вектора на число. Скалярное произведение векторов"
)
NEW_TOPIC = "Первообразная и определённый интеграл"


@pytest.fixture
def local_project(tmp_path, monkeypatch):
    missing = [name for name in SOURCE_NAMES if not (PROJECT / "input" / name).is_file()]
    if missing:
        pytest.skip("Для проверки нужны сохранённые реальные исходники: " + ", ".join(missing))
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    for name in SOURCE_NAMES:
        shutil.copyfile(PROJECT / "input" / name, input_dir / name)
    shutil.copytree(PROJECT / "data", tmp_path / "data")
    monkeypatch.setattr(main, "BASE_DIR", tmp_path)
    monkeypatch.setattr(main, "INPUT_DIR", input_dir)
    monkeypatch.delenv("GEMINI_API_KEY", raising=False)

    def network_forbidden(*args, **kwargs):
        pytest.fail("Сеть при локальной сборке запрещена")

    monkeypatch.setattr(socket.socket, "connect", network_forbidden)
    monkeypatch.setattr(downloader, "download_bytes", network_forbidden)
    monkeypatch.setattr(recognizer.genai, "Client", network_forbidden)
    return tmp_path


def fake_classification(prompt, *, model, schema):
    """Данные для проверки контракта, не математическая классификация варианта."""
    payload = json.loads(prompt.split("Данные:\n", 1)[1])
    assert len(payload["tasks"]) == 19
    assert EXAM_LABEL in payload["exams"]
    assert "E | Вектор" not in payload["exams"]
    return json.dumps([
        {
            "task_num": number,
            "exam_choice": EXAM_LABEL,
            "topic_kind": "new" if number == 8 else "existing",
            "topic_choice": NEW_TOPIC if number == 8 else TOPIC_LABEL,
        }
        for number in range(1, 20)
    ], ensure_ascii=False)


def test_real_sources_and_both_caches_rebuild_with_network_blocked(local_project, monkeypatch):
    calls = []

    def classify_once(*args, **kwargs):
        calls.append(args)
        return fake_classification(*args, **kwargs)

    original_sources = {name: (local_project / "input" / name).read_bytes() for name in SOURCE_NAMES}
    monkeypatch.setattr(classification, "request_gemini", classify_once)
    main.main(main.DEFAULT_PAGE_URL)
    assert len(calls) == 1

    def gemini_forbidden(*args, **kwargs):
        pytest.fail("Gemini при подходящем кэше запрещён")

    monkeypatch.setattr(classification, "request_gemini", gemini_forbidden)
    main.main(main.DEFAULT_PAGE_URL)
    for name, contents in original_sources.items():
        assert (local_project / "input" / name).read_bytes() == contents

    output = local_project / "output" / "trvar542"
    expected = {"tasks.xlsx", *IMAGE_NAMES.values()}
    assert {path.name for path in output.iterdir()} == expected
    assert not list(output.glob("*ocr*"))
    with ZipFile(local_project / "trvar542_result.zip") as archive:
        assert set(archive.namelist()) == expected
        assert len(archive.namelist()) == 5
        assert archive.testzip() is None
        for name in expected:
            assert archive.read(name) == (output / name).read_bytes()

    workbook = load_workbook(output / "tasks.xlsx")
    try:
        sheet = workbook["Tasks"]
        assert list(next(sheet.values)) == ["task_num", "condition", "image_name", "solution", "answer", "exams_id", "topics_id"]
        assert sheet.max_row == 20
        assert [sheet.cell(row, 1).value for row in range(2, 21)] == list(range(1, 20))
        assert {row[0].value: row[2].value for row in sheet.iter_rows(min_row=2) if row[2].value} == IMAGE_NAMES
        assert sheet["E6"].value == "0,9940"
        assert sheet["E6"].data_type == "s"
        assert sheet["G9"].value == NEW_TOPIC
        assert sheet["G9"].data_type == "s"
        assert sheet["G3"].value == 526
        assert sheet["G3"].data_type == "n"
        assert all(row[3].value is None for row in sheet.iter_rows(min_row=2))
        assert all(type(row[5].value) is int for row in sheet.iter_rows(min_row=2))
        # Проверяем конечный Excel по полному исходному справочнику,
        # независимо от функции отбора кандидатов классификатора.
        with (local_project / "data" / "exams.csv").open(encoding="utf-8-sig") as source:
            exams = {int(record["id"]): int(record["parent_of_exam"]) for record in csv.DictReader(source)}
        parents = set(exams.values())
        for row in sheet.iter_rows(min_row=2):
            exam_id = row[5].value
            assert exam_id == 180
            assert exam_id not in parents
            current = exams[exam_id]
            while current and current != 2:
                current = exams[current]
            assert current == 2
        assert "<br>Б)" in sheet["E14"].value
        assert "\\frac" in sheet["E14"].value
        assert "<br>Б)" in sheet["B14"].value
        assert "<br>Б)" in sheet["E20"].value
        about = workbook["about"]
        assert list(next(about.values)) == list(ABOUT_FIELDS)
        assert about.max_column == 11 and about.max_row == 2
        fields = dict(zip((cell.value for cell in about[1]), (cell.value for cell in about[2])))
        assert fields["source_url"] == main.DEFAULT_PAGE_URL
        assert fields["exam_id"] is None
    finally:
        workbook.close()


def test_changed_prompt_reclassifies_topics_through_main_and_reuses_both_caches(
    local_project, monkeypatch,
):
    """Проверяем обновление выгрузки через main; смысловой выбор API подменён."""
    original_sources = {
        name: (local_project / "input" / name).read_bytes() for name in SOURCE_NAMES
    }
    topics = classification.load_catalog(
        local_project / "data" / "topics.csv",
        name_column="name", parent_column="parent", root_id=1,
    )
    labels = {
        number: label
        for label, number in classification.build_category_choices(topics, "T").items()
    }
    previous_prompt = "Предыдущая инструкция классификации без смысловых уточнений."
    current_prompt = classification.PROMPT
    calls = []
    assignments = {3: 616, 11: 541, 19: 643}

    def fake_response(prompt, *, model, schema):
        calls.append(prompt)
        items = json.loads(fake_classification(prompt, model=model, schema=schema))
        for item in items:
            if item["task_num"] not in assignments:
                continue
            topic = assignments[item["task_num"]]
            item["topic_kind"] = "existing" if type(topic) is int else "new"
            item["topic_choice"] = labels[topic] if type(topic) is int else topic
        return json.dumps(items, ensure_ascii=False)

    monkeypatch.setattr(classification, "request_gemini", fake_response)
    monkeypatch.setattr(classification, "PROMPT", previous_prompt)
    main.main(main.DEFAULT_PAGE_URL)
    cache_path = local_project / "input" / "trvar542_classification.json"
    previous_cache = json.loads(cache_path.read_text(encoding="utf-8"))
    assert {row["task_num"]: row["topics_id"] for row in previous_cache["tasks"]
            if row["task_num"] in assignments} == assignments

    # Меняется только инструкция: правила структуры, данные и настройки те же.
    monkeypatch.setattr(classification, "PROMPT", current_prompt)
    assignments = {3: 585, 11: "Кусочно заданная функция", 19: 637}
    main.main(main.DEFAULT_PAGE_URL)
    current_cache = json.loads(cache_path.read_text(encoding="utf-8"))
    assert len(calls) == 2
    assert calls[0].startswith(previous_prompt)
    assert calls[1].startswith(current_prompt)
    assert current_cache["rules_version"] == previous_cache["rules_version"]
    assert current_cache["fingerprint"] != previous_cache["fingerprint"]
    assert current_cache["choices"][10]["topic_kind"] == "new"
    assert {row["task_num"]: row["topics_id"] for row in current_cache["tasks"]
            if row["task_num"] in assignments} == assignments

    def gemini_forbidden(*args, **kwargs):
        pytest.fail("Повторная сборка с актуальным промптом должна читать кэш")

    monkeypatch.setattr(classification, "request_gemini", gemini_forbidden)
    main.main(main.DEFAULT_PAGE_URL)
    assert json.loads(cache_path.read_text(encoding="utf-8")) == current_cache
    assert all((local_project / "input" / name).read_bytes() == contents
               for name, contents in original_sources.items())

    workbook = load_workbook(local_project / "output" / "trvar542" / "tasks.xlsx")
    try:
        sheet = workbook["Tasks"]
        assert sheet.max_column == 7 and sheet.max_row == 20
        assert sheet["G4"].value == 585 and sheet["G4"].data_type == "n"
        assert sheet["G12"].value == "Кусочно заданная функция"
        assert sheet["G12"].data_type == "s"
        assert sheet["G20"].value == 637 and sheet["G20"].data_type == "n"
    finally:
        workbook.close()


@pytest.mark.parametrize("failure", ["api", "invalid_response"])
def test_failed_classification_keeps_previous_successful_files(local_project, monkeypatch, failure):
    monkeypatch.setattr(classification, "request_gemini", fake_classification)
    main.main(main.DEFAULT_PAGE_URL)
    output = local_project / "output" / "trvar542"
    before = {path.name: path.read_bytes() for path in output.iterdir()}
    archive = local_project / "trvar542_result.zip"
    original_archive = archive.read_bytes()
    (local_project / "input" / "trvar542_classification.json").unlink()

    calls = []

    def fail_classification(*args, **kwargs):
        calls.append(args)
        if failure == "api":
            raise RuntimeError("Gemini: сеть или API недоступны")
        return "Некорректный JSON"

    monkeypatch.setattr(classification, "request_gemini", fail_classification)
    error_type, message = (RuntimeError, "Gemini") if failure == "api" else (ValueError, "после двух запросов")
    with pytest.raises(error_type, match=message):
        main.main(main.DEFAULT_PAGE_URL)
    assert len(calls) == (1 if failure == "api" else 2)
    assert {path.name: path.read_bytes() for path in output.iterdir()} == before
    assert archive.read_bytes() == original_archive
    assert not (local_project / "input" / "previous_output").exists()


@pytest.mark.parametrize("encoding", ["utf-8", "windows-1251"])
def test_html_uses_original_url_and_exact_variant(encoding):
    html = '''<html><title>Тренировочный вариант №542 ЕГЭ - 2026</title>
    <a href="other.pdf">Other</a><a href="trvar542.pdf">PDF</a>
    <script src="%74rvar542.js"></script><img src="trvar542.png"></html>'''
    assert main.parse_source_page(html.encode(encoding), main.DEFAULT_PAGE_URL) == {
        extension: f"https://alexlarin.net/ege/2026/trvar542.{extension}"
        for extension in ("pdf", "js", "png")
    }
    with pytest.raises(ValueError, match="нужного тренировочного варианта"):
        main.parse_source_page(html.replace("№542", "№541").encode(encoding), main.DEFAULT_PAGE_URL)
    with pytest.raises(ValueError, match="Год"):
        main.parse_source_page(html.replace("2026", "2025").encode(encoding), main.DEFAULT_PAGE_URL)
    with pytest.raises(ValueError, match="trvar542.pdf"):
        main.parse_source_page(html.replace("trvar542.pdf", "trvar541.pdf").encode(encoding), main.DEFAULT_PAGE_URL)


@pytest.mark.parametrize("reader", [main.read_js_source, main.validate_answer_image])
def test_error_page_cannot_replace_js_or_image(reader):
    with pytest.raises(ValueError):
        reader(b"<html>503 Service Unavailable</html>")


def test_corrupt_packed_js_is_rejected_with_clear_error():
    with pytest.raises(ValueError, match="JavaScript"):
        main.read_js_source(b"eval(function(p,a,c,k,e,r){broken")


def test_real_pdf_with_wrong_number_or_year_is_rejected(local_project):
    data = (local_project / "input" / "trvar542.pdf").read_bytes()
    with pytest.raises(ValueError, match="другому варианту"):
        main.read_pdf_source(data, "https://alexlarin.net/ege/2026/trvar541.html")
    with pytest.raises(ValueError, match="Год"):
        main.read_pdf_source(data, "https://alexlarin.net/ege/2025/trvar542.html")


@pytest.mark.parametrize("extension", ["pdf", "js", "png"])
@pytest.mark.parametrize("wrong_prefix", [
    "https://other.example/ege/2026/", "https://alexlarin.net/ege/2025/",
])
def test_html_rejects_resource_from_other_origin_or_year(extension, wrong_prefix):
    html = '''<html><title>Тренировочный вариант №542 ЕГЭ - 2026</title>
    <a href="trvar542.pdf">PDF</a><script src="trvar542.js"></script>
    <img src="trvar542.png"></html>'''
    wrong_html = html.replace(f'"trvar542.{extension}"', f'"{wrong_prefix}trvar542.{extension}"')
    with pytest.raises(ValueError, match=f"trvar542.{extension}"):
        main.parse_source_page(wrong_html.encode(), main.DEFAULT_PAGE_URL)


@pytest.fixture
def extra_answers():
    return {
        "source_url": "https://alexlarin.net/ege/2026/trvar542.png",
        "method": "manual_transcription",
        "answers": {str(number): "0,9940" if number == 13 else "$\\frac{2}{3}$" for number in range(13, 20)},
    }


def test_extra_answers_keep_original_strings_and_validate_source(tmp_path, extra_answers):
    path = tmp_path / "answers.json"
    path.write_text(json.dumps(extra_answers), encoding="utf-8")
    result = main.read_extra_answers(path, extra_answers["source_url"])
    assert list(result["answers"]) == list(range(13, 20))
    assert result["answers"][13] == "0,9940"
    assert result["answers"][17] == "$\\frac{2}{3}$"
    assert result["method"] == "manual_transcription"
    with pytest.raises(ValueError, match="другой картинке"):
        main.read_extra_answers(path, "https://alexlarin.net/ege/2026/trvar541.png")


@pytest.mark.parametrize("corruption", [
    "missing", "invalid_json", "not_object", "missing_answers", "not_answers_object",
    "missing_method", "invalid_number", "empty_answer", "boolean_answer", "missing_task", "duplicate_task",
])
def test_extra_answers_fail_with_clear_error(tmp_path, extra_answers, corruption):
    path = tmp_path / "answers.json"
    image_url = extra_answers["source_url"]
    if corruption == "not_object":
        extra_answers = []
    elif corruption == "missing_answers":
        extra_answers.pop("answers")
    elif corruption == "not_answers_object":
        extra_answers["answers"] = []
    elif corruption == "missing_method":
        extra_answers.pop("method")
    elif corruption == "invalid_number":
        extra_answers["answers"]["thirteen"] = extra_answers["answers"].pop("13")
    elif corruption == "empty_answer":
        extra_answers["answers"]["13"] = " "
    elif corruption == "boolean_answer":
        extra_answers["answers"]["13"] = True
    elif corruption == "missing_task":
        extra_answers["answers"].pop("13")
    elif corruption == "duplicate_task":
        extra_answers["answers"]["013"] = "1"
    if corruption != "missing":
        path.write_text("broken JSON" if corruption == "invalid_json" else json.dumps(extra_answers), encoding="utf-8")
    with pytest.raises(ValueError, match="answers.json"):
        main.read_extra_answers(path, image_url)

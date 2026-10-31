"""Проверки структуры Excel и расположения файлов в итоговом архиве."""

from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

from openpyxl import load_workbook

from export_metadata import ABOUT_FIELDS, prepare_about_metadata
from exporter import create_zip, write_excel


class ExporterTests(unittest.TestCase):
    def setUp(self):
        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.output = self.root / "output" / "trvar542"
        self.output.mkdir(parents=True)
        self.excel_path = self.output / "tasks.xlsx"
        self.image_names = {
            1: "image_82.jpeg",
            3: "image_83.jpeg",
            8: "image_198.png",
            11: "image_200.png",
        }
        self.tasks = [
            {
                "task_num": number,
                "condition": r"Найдите длину $AA_1$ и значение $\sqrt{43}$.",
                "image_name": self.image_names.get(number, ""),
                "answer": "0,9940" if number == 5 else "12",
                "exams_id": 179,
                "topics_id": 526,
            }
            for number in range(1, 20)
        ]
        self.metadata = {
            "class": 11,
            "year": 2026,
            "date": date(2026, 10, 20),
            "variant_name": "Тренировочный вариант № 542",
            "source_url": "https://alexlarin.net/ege/2026/trvar542.html",
            "conditions_format": "LaTeX для отображения через MathJax",
        }

    def _write_result(self):
        write_excel(self.tasks, self.metadata, self.excel_path)
        for name in self.image_names.values():
            (self.output / name).write_bytes(f"test image: {name}".encode())

    def _assert_archive(self, archive_path):
        expected = {"tasks.xlsx", *self.image_names.values()}
        with ZipFile(archive_path) as archive:
            self.assertEqual(set(archive.namelist()), expected)
            self.assertEqual(len(archive.namelist()), 5)
            self.assertIsNone(archive.testzip())
            for name in expected:
                self.assertEqual(archive.read(name), (self.output / name).read_bytes())

    def test_about_has_headers_and_values_in_two_rows(self):
        write_excel(self.tasks, self.metadata, self.excel_path)
        workbook = load_workbook(self.excel_path)
        self.addCleanup(workbook.close)
        about = workbook["about"]
        self.assertEqual(workbook.sheetnames, ["Tasks", "about"])
        self.assertEqual(about.max_row, 2)
        self.assertEqual([cell.value for cell in about[1]], list(self.metadata))
        self.assertEqual(about["A2"].value, 11)
        self.assertEqual(about["B2"].value, 2026)
        self.assertEqual(about["C2"].value.date(), self.metadata["date"])
        self.assertEqual(about["C2"].number_format, "yyyy-mm-dd")
        self.assertEqual(about["D2"].value, self.metadata["variant_name"])
        self.assertEqual(about["E2"].value, self.metadata["source_url"])
        self.assertEqual(about["F2"].value, self.metadata["conditions_format"])

    def test_prepared_about_preserves_eleven_fields_and_source(self):
        metadata = prepare_about_metadata(self.metadata)
        write_excel(self.tasks, metadata, self.excel_path)
        workbook = load_workbook(self.excel_path)
        self.addCleanup(workbook.close)
        about = workbook["about"]
        self.assertEqual([cell.value for cell in about[1]], list(ABOUT_FIELDS))
        self.assertEqual(about.max_column, 11)
        self.assertEqual(about.max_row, 2)
        self.assertEqual(about["F2"].value, self.metadata["variant_name"])
        self.assertEqual(about["H2"].value, "alexlarin.net")
        self.assertEqual(about["J2"].value, self.metadata["source_url"])
        self.assertIsNone(about["E2"].value)

    def test_tasks_preserve_latex_and_text_answers(self):
        write_excel(self.tasks, self.metadata, self.excel_path)
        workbook = load_workbook(self.excel_path)
        self.addCleanup(workbook.close)
        tasks = workbook["Tasks"]
        self.assertEqual(
            [cell.value for cell in tasks[1]],
            ["task_num", "condition", "image_name", "solution", "answer", "exams_id", "topics_id"],
        )
        self.assertEqual(tasks.max_row, 20)
        for row, task in zip(tasks.iter_rows(min_row=2), self.tasks):
            self.assertEqual(row[0].value, task["task_num"])
            self.assertEqual(row[1].value, task["condition"])
            self.assertEqual(row[2].value or "", task["image_name"])
            self.assertIsNone(row[3].value)
            self.assertEqual(row[4].value, task["answer"])
            self.assertEqual(row[4].data_type, "s")
            self.assertEqual(row[5].value, task["exams_id"])
            self.assertEqual(row[6].value, task["topics_id"])
            self.assertEqual(row[5].data_type, "n")
            self.assertEqual(row[6].data_type, "n")
        self.assertEqual(tasks["E6"].value, "0,9940")
        self.assertTrue(tasks["B2"].alignment.wrap_text)
        self.assertEqual(tasks.freeze_panes, "A2")
        self.assertEqual(tasks.auto_filter.ref, "A1:G20")

    def test_mixed_topic_types_stay_in_seven_columns(self):
        self.tasks[7]["topics_id"] = "Первообразная и определённый интеграл"
        self.tasks[12]["condition"] = r"А) Найдите $x$.<br>Б) Найдите $\sqrt{x}$."
        self.tasks[12]["answer"] = r"А) $2$.<br>Б) $\sqrt{2}$."
        write_excel(self.tasks, prepare_about_metadata(self.metadata), self.excel_path)
        workbook = load_workbook(self.excel_path)
        self.addCleanup(workbook.close)
        tasks = workbook["Tasks"]

        self.assertEqual(tasks.max_column, 7)
        self.assertEqual(tasks.max_row, 20)
        self.assertEqual([row[0].value for row in tasks.iter_rows(min_row=2)], list(range(1, 20)))
        for row, task in zip(tasks.iter_rows(min_row=2), self.tasks):
            self.assertEqual(row[2].value or "", self.image_names.get(task["task_num"], ""))
            self.assertIsNone(row[3].value)
            self.assertEqual(row[4].value, task["answer"])
            self.assertEqual(row[4].data_type, "s")
            self.assertEqual(row[5].value, 179)
            self.assertEqual(row[5].data_type, "n")
            self.assertEqual(row[6].value, task["topics_id"])
            self.assertEqual(row[6].data_type, "s" if task["task_num"] == 8 else "n")
        self.assertEqual(tasks["B14"].value, self.tasks[12]["condition"])
        self.assertEqual(tasks["E6"].value, "0,9940")

    def test_text_starting_with_equals_is_not_an_excel_formula(self):
        self.tasks[7]["topics_id"] = "=Первообразная и определённый интеграл"
        self.tasks[0]["answer"] = "=2+2"
        write_excel(self.tasks, self.metadata, self.excel_path)
        for data_only in (False, True):
            with self.subTest(data_only=data_only):
                workbook = load_workbook(self.excel_path, data_only=data_only)
                try:
                    tasks = workbook["Tasks"]
                    self.assertEqual(tasks["G9"].value, self.tasks[7]["topics_id"])
                    self.assertEqual(tasks["G9"].data_type, "s")
                    self.assertEqual(tasks["E2"].value, "=2+2")
                    self.assertEqual(tasks["E2"].data_type, "s")
                finally:
                    workbook.close()

    def test_zip_has_five_files_at_one_level(self):
        self._write_result()
        archive_path = self.root / "trvar542_result.zip"
        count = create_zip(self.excel_path, self.output, archive_path)
        self.assertEqual(count, 5)
        self._assert_archive(archive_path)

    def test_repeated_zip_does_not_include_itself(self):
        self._write_result()
        archive_path = self.output / "result.zip"
        for _ in range(2):
            self.assertEqual(create_zip(self.excel_path, self.output, archive_path), 5)
            self._assert_archive(archive_path)

    def test_zip_also_accepts_a_separate_image_directory(self):
        self._write_result()
        images = self.root / "images"
        images.mkdir()
        for name in self.image_names.values():
            (images / name).write_bytes((self.output / name).read_bytes())
        archive_path = self.root / "separate.zip"
        self.assertEqual(create_zip(self.excel_path, images, archive_path), 5)
        self._assert_archive(archive_path)


if __name__ == "__main__":
    unittest.main()
